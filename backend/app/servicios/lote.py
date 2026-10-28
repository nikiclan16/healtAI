"""Procesamiento por lotes (RF03.2): CSV/XLSX → inferencia + reglas por fila, sin persistir.

Las columnas del archivo son las mismas variables del formulario (sin identificadores del paciente).
Una columna desconocida se rechaza: así un identificador (cédula, nombre…) o el nivel de riesgo de
amputación no entran por esta vía. Cada fila se valida con el mismo contrato que `/evaluaciones`.
"""

import csv
import hashlib
import io
from typing import Any

from fastapi import HTTPException
from pydantic import ValidationError

from .. import APP_VERSION
from ..config import get_settings
from ..modelo.component import ComponenteEstratificacionSaludMental
from ..schemas_clinicos import ScreeningFormIn
from . import evaluacion as svc

MAX_BYTES = 5 * 1024 * 1024
MAX_FILAS = 2000

# columna del archivo → (ruta en el formulario, tipo)
_ENTERO, _BOOL, _TRAT = "entero", "bool", "tratamiento"
COLUMNAS: dict[str, tuple[tuple[str, ...], str]] = {
    "v8_tiempo_diagnostico": (("v8_tiempo_diagnostico",), _ENTERO),
    "v9_tratamiento": (("v9_tratamiento",), _TRAT),
    "v11_hipertension": (("v11", "hipertension"), _BOOL),
    "v11_cardiaca": (("v11", "cardiaca"), _BOOL),
    "v11_obesidad": (("v11", "obesidad"), _BOOL),
    "v11_otra": (("v11", "otra"), _BOOL),
    "v12_percepcion_salud": (("v12_percepcion_salud",), _ENTERO),
    "v13_apoyo_familiar": (("v13_apoyo_familiar",), _ENTERO),
    "v14_actividad_fisica": (("v14_actividad_fisica",), _ENTERO),
    "v15_calidad_sueno": (("v15_calidad_sueno",), _ENTERO),
    "who5_total": (("who5_total",), _ENTERO),
    "dass_depresion": (("dass_depresion",), _ENTERO),
    "dass_ansiedad": (("dass_ansiedad",), _ENTERO),
    "dass_estres": (("dass_estres",), _ENTERO),
    "dass_item21": (("dass_item21",), "item21"),
    # Opcionales: si faltan, se asumen sin alerta del profesional.
    "alerta_riesgo_suicida": (("alertas_profesional", "riesgo_suicida"), "alerta"),
    "alerta_deterioro_agudo": (("alertas_profesional", "deterioro_agudo"), "alerta"),
    "alerta_riesgo_urgente": (("alertas_profesional", "riesgo_urgente"), "alerta"),
}
OPCIONALES = {c for c, (_, tipo) in COLUMNAS.items() if tipo == "alerta"}
REQUERIDAS = [c for c in COLUMNAS if c not in OPCIONALES]

_SI = {"1", "true", "si", "sí", "s", "yes", "y", "x", "verdadero"}
_NO = {"0", "false", "no", "n", "falso"}
_SIN_DATO = {"", "na", "n/a", "nan", "none", "null", "no evaluable", "no_evaluable"}

PLANTILLA = [
    ",".join(COLUMNAS),
    "10,orales,1,0,0,0,3,4,0,2,8,9,8,12,1,0,0,0",
]


class FilaInvalida(ValueError):
    pass


def _texto(valor: Any) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _convertir(columna: str, tipo: str, valor: Any) -> Any:
    t = _texto(valor)
    k = t.lower()
    if tipo == _ENTERO:
        try:
            f = float(t.replace(",", "."))
        except ValueError:
            raise FilaInvalida(f"{columna}: se esperaba un número entero") from None
        if not f.is_integer():
            raise FilaInvalida(f"{columna}: se esperaba un número entero")
        return int(f)
    if tipo == "item21":
        return None if k in _SIN_DATO else _convertir(columna, _ENTERO, valor)
    if tipo in (_BOOL, "alerta"):
        if k in _SI:
            return True
        if k in _NO or (tipo == "alerta" and k == ""):
            return False
        raise FilaInvalida(f"{columna}: use 1/0, sí/no o true/false")
    return k  # tratamiento: lo valida el contrato


def _mensaje(error: dict[str, Any]) -> str:
    """Mensaje de validación en español para el usuario del lote."""
    ctx = error.get("ctx") or {}
    match error["type"]:
        case "less_than_equal":
            return f"debe ser como máximo {ctx.get('le')}"
        case "greater_than_equal":
            return f"debe ser como mínimo {ctx.get('ge')}"
        case "literal_error":
            return f"valor no permitido (use: {ctx.get('expected', '').replace(chr(39), '')})"
        case "missing":
            return "dato obligatorio"
        case _:
            return str(error["msg"])


def _formulario(fila: dict[str, Any]) -> ScreeningFormIn:
    datos: dict[str, Any] = {"alertas_profesional": {"riesgo_suicida": False, "deterioro_agudo": False, "riesgo_urgente": False}, "v11": {}}
    for columna, (ruta, tipo) in COLUMNAS.items():
        if columna not in fila:
            continue  # solo las opcionales pueden faltar (se validó la cabecera)
        valor = _convertir(columna, tipo, fila[columna])
        destino = datos
        for parte in ruta[:-1]:
            destino = destino.setdefault(parte, {})
        destino[ruta[-1]] = valor
    try:
        return ScreeningFormIn.model_validate(datos)
    except ValidationError as exc:
        inverso = {ruta: col for col, (ruta, _) in COLUMNAS.items()}
        errores = []
        for e in exc.errors():
            loc = tuple(str(p) for p in e["loc"])
            errores.append(f"{inverso.get(loc, '.'.join(loc))}: {_mensaje(e)}")
        raise FilaInvalida("; ".join(errores)) from None


# ── Lectura del archivo ─────────────────────────────────────────────────────────

def _leer_csv(contenido: bytes) -> list[list[Any]]:
    try:
        texto = contenido.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = contenido.decode("latin-1")
    muestra = texto[:4096]
    delimitador = max(",;\t", key=muestra.count)  # Excel en español exporta con ';'
    return [list(r) for r in csv.reader(io.StringIO(texto), delimiter=delimitador)]


def _leer_xlsx(contenido: bytes) -> list[list[Any]]:
    from openpyxl import load_workbook

    try:
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(422, "No se pudo leer el archivo XLSX") from None
    try:
        return [list(r) for r in libro.worksheets[0].iter_rows(values_only=True)]
    finally:
        libro.close()


def leer_tabla(nombre: str, contenido: bytes) -> tuple[list[str], list[tuple[int, dict[str, Any]]]]:
    """Devuelve (cabecera normalizada, [(número de fila en el archivo, valores por columna)])."""
    if len(contenido) > MAX_BYTES:
        raise HTTPException(413, f"El archivo supera el máximo de {MAX_BYTES // (1024 * 1024)} MB")
    extension = nombre.lower().rsplit(".", 1)[-1] if "." in nombre else ""
    if extension == "xlsx":
        filas = _leer_xlsx(contenido)
    elif extension in ("csv", "txt"):
        filas = _leer_csv(contenido)
    else:
        raise HTTPException(415, "Formato no admitido: use un archivo .csv o .xlsx")

    filas = [r for r in filas if any(_texto(c) for c in r)]  # descarta filas vacías
    if not filas:
        raise HTTPException(422, "El archivo está vacío")
    cabecera = [_texto(c).lower() for c in filas[0]]
    if len(filas) - 1 > MAX_FILAS:
        raise HTTPException(413, f"El archivo supera el máximo de {MAX_FILAS} filas por lote")

    desconocidas = [c for c in cabecera if c and c not in COLUMNAS]
    if desconocidas:
        raise HTTPException(
            422,
            "Columnas no reconocidas (no se admiten identificadores del paciente ni columnas fuera del "
            f"Diccionario {get_settings().dictionary_version}): {', '.join(desconocidas)}",
        )
    repetidas = sorted({c for c in cabecera if c and cabecera.count(c) > 1})
    if repetidas:
        raise HTTPException(422, f"Columnas repetidas: {', '.join(repetidas)}")
    faltantes = [c for c in REQUERIDAS if c not in cabecera]
    if faltantes:
        raise HTTPException(422, f"Faltan columnas obligatorias: {', '.join(faltantes)}")

    registros = []
    for numero, fila in enumerate(filas[1:], start=2):
        registros.append((numero, {c: fila[i] if i < len(fila) else None for i, c in enumerate(cabecera) if c}))
    return cabecera, registros


# ── Procesamiento ───────────────────────────────────────────────────────────────

def procesar(modelo: ComponenteEstratificacionSaludMental, nombre: str, contenido: bytes) -> dict:
    cabecera, registros = leer_tabla(nombre, contenido)
    resultados = []
    for numero, fila in registros:
        try:
            r = svc.ejecutar(modelo, _formulario(fila))
        except FilaInvalida as exc:
            resultados.append({"fila": numero, "ok": False, "error": str(exc)})
            continue
        resultados.append(
            {
                "fila": numero,
                "ok": True,
                "probabilidadRA": r["prediccion"]["probabilidadRA"],
                "rutaBase": r["prediccion"]["rutaBase"],
                "rutaSistema": r["rutaSistema"],
                "escalada": r["escalada"],
                "reglas": [x["code"] for x in r["reglas"]["triggered"]],
                "shapTop3": r["prediccion"]["shapTop3"],
            }
        )

    ok = [x for x in resultados if x["ok"]]
    return {
        "archivo": nombre[:120],
        "sha256": hashlib.sha256(contenido).hexdigest(),
        "columnas": [c for c in cabecera if c],
        "resumen": {
            "filas": len(resultados),
            "procesadas": len(ok),
            "conError": len(resultados) - len(ok),
            "baseRA": sum(x["rutaBase"] == "A" for x in ok),
            "baseRB": sum(x["rutaBase"] == "B" for x in ok),
            "escaladas": sum(x["escalada"] for x in ok),
            "finalA": sum(x["rutaSistema"] == "A" for x in ok),
            "finalB": sum(x["rutaSistema"] == "B" for x in ok),
        },
        "versiones": {
            "app": APP_VERSION,
            "modelo": modelo.artefacto.version,
            "diccionario": get_settings().dictionary_version,
        },
        "modeloPlaceholder": modelo.artefacto.placeholder,
        "resultados": resultados,
    }
